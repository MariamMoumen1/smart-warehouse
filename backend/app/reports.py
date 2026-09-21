from io import BytesIO
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def generate_pdf_report(status: dict, trucks: list, predictions: list) -> BytesIO:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleCustom', parent=styles['Title'], textColor=colors.HexColor('#4527a0'))
    heading_style = ParagraphStyle('HeadingCustom', parent=styles['Heading2'], textColor=colors.HexColor('#16213e'), spaceBefore=16, spaceAfter=8)

    elements = []

    elements.append(Paragraph("Smart Warehouse — Rapport Global", title_style))
    elements.append(Paragraph(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 20))

    # État de l'entrepôt
    elements.append(Paragraph("État de l'entrepôt", heading_style))
    status_data = [
        ["Niveau de stock", f"{status['niveau_stock']} / {status['capacite_max']}"],
        ["Camions en cours", str(status['camions_entrants'])],
        ["Camions en attente", str(status['camions_en_attente'])],
    ]
    t = Table(status_data, colWidths=[8*cm, 8*cm])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f4f6fb')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e0e4ef')),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('PADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(t)

    # Alertes
    elements.append(Paragraph("Alertes actives", heading_style))
    if status['alertes']:
        for alerte in status['alertes']:
            elements.append(Paragraph(f"• [{alerte['type'].upper()}] {alerte['message']}", styles['Normal']))
    else:
        elements.append(Paragraph("Aucune alerte active", styles['Normal']))

    # Camions
    elements.append(Paragraph(f"Camions ({len(trucks)})", heading_style))
    if trucks:
        truck_data = [["ID", "Plaque", "Statut", "Créé le"]]
        for tr in trucks:
            truck_data.append([str(tr.id), tr.plaque, tr.statut, tr.created_at.strftime('%d/%m/%Y %H:%M')])
        t2 = Table(truck_data, colWidths=[2*cm, 5*cm, 4*cm, 5*cm])
        t2.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4527a0')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e0e4ef')),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('PADDING', (0, 0), (-1, -1), 6),
        ]))
        elements.append(t2)
    else:
        elements.append(Paragraph("Aucun camion enregistré", styles['Normal']))

    # Prédictions
    elements.append(Paragraph(f"Historique des prédictions ({len(predictions)})", heading_style))
    if predictions:
        pred_data = [["Date", "Heure", "Camions", "Stock", "Résultat (min)"]]
        for p in predictions[:20]:
            pred_data.append([
                p.created_at.strftime('%d/%m %H:%M'),
                f"{p.heure}h",
                str(p.camions_entrants),
                str(p.niveau_stock),
                str(p.resultat)
            ])
        t3 = Table(pred_data, colWidths=[3.5*cm, 2.5*cm, 3*cm, 3*cm, 4*cm])
        t3.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4527a0')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e0e4ef')),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('PADDING', (0, 0), (-1, -1), 6),
        ]))
        elements.append(t3)
    else:
        elements.append(Paragraph("Aucune prédiction enregistrée", styles['Normal']))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generate_excel_report(status: dict, trucks: list, predictions: list) -> BytesIO:
    wb = Workbook()

    header_fill = PatternFill(start_color="4527A0", end_color="4527A0", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    # Feuille 1 : État de l'entrepôt
    ws1 = wb.active
    ws1.title = "État entrepôt"
    ws1.append(["Smart Warehouse — Rapport Global"])
    ws1.append([f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"])
    ws1.append([])
    ws1.append(["Indicateur", "Valeur"])
    ws1["A4"].fill = header_fill
    ws1["B4"].fill = header_fill
    ws1["A4"].font = header_font
    ws1["B4"].font = header_font
    ws1.append(["Niveau de stock", f"{status['niveau_stock']} / {status['capacite_max']}"])
    ws1.append(["Camions en cours", status['camions_entrants']])
    ws1.append(["Camions en attente", status['camions_en_attente']])
    ws1.append([])
    ws1.append(["Alertes"])
    for alerte in status['alertes']:
        ws1.append([f"[{alerte['type'].upper()}] {alerte['message']}"])
    ws1.column_dimensions['A'].width = 30
    ws1.column_dimensions['B'].width = 30

    # Feuille 2 : Camions
    ws2 = wb.create_sheet("Camions")
    ws2.append(["ID", "Plaque", "Statut", "Temps attente (min)", "Créé le"])
    for col in ['A', 'B', 'C', 'D', 'E']:
        ws2[f"{col}1"].fill = header_fill
        ws2[f"{col}1"].font = header_font
    for tr in trucks:
        ws2.append([tr.id, tr.plaque, tr.statut, tr.temps_attente_min, tr.created_at.strftime('%d/%m/%Y %H:%M')])
    for col, width in zip(['A', 'B', 'C', 'D', 'E'], [8, 18, 15, 18, 20]):
        ws2.column_dimensions[col].width = width

    # Feuille 3 : Prédictions
    ws3 = wb.create_sheet("Prédictions")
    ws3.append(["Date", "Heure", "Jour semaine", "Camions entrants", "Stock", "Température", "Résultat (min)"])
    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
        ws3[f"{col}1"].fill = header_fill
        ws3[f"{col}1"].font = header_font
    for p in predictions:
        ws3.append([
            p.created_at.strftime('%d/%m/%Y %H:%M'), p.heure, p.jour_semaine,
            p.camions_entrants, p.niveau_stock, p.temperature, p.resultat
        ])
    for col, width in zip(['A', 'B', 'C', 'D', 'E', 'F', 'G'], [18, 8, 12, 15, 10, 12, 15]):
        ws3.column_dimensions[col].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
